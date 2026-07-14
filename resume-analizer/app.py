import streamlit as st
import fitz  # PyMuPDF
import docx
import spacy
from spacy.matcher import PhraseMatcher
from datetime import datetime
import json

import database
import pdf_generator

# -----------------------------
# Skill Database
# -----------------------------
COMMON_SKILLS = [
    "python", "java", "c", "c++", "javascript",
    "html", "css", "react", "angular", "nodejs",
    "django", "flask", "sql", "mysql", "mongodb",
    "postgresql", "git", "github", "docker",
    "kubernetes", "aws", "azure", "machine learning",
    "data analysis", "power bi", "excel",
    "rest api", "spring boot", "typescript", "c#",
    "go", "rust", "devops", "ci/cd", "tableau",
    "pandas", "numpy", "scikit-learn", "tensorflow", "pytorch"
]

# Initialize Database
database.init_db()

# -----------------------------
# Load spaCy Model
# -----------------------------
@st.cache_resource
def load_spacy_model():
    try:
        return spacy.load("en_core_web_sm")
    except OSError:
        with st.spinner("Downloading NLP language model (en_core_web_sm)... This may take a moment."):
            from spacy.cli import download
            download("en_core_web_sm")
            return spacy.load("en_core_web_sm")

# -----------------------------
# Extract Text From PDF
# -----------------------------
def extract_text_from_pdf(pdf_file):
    pdf_file.seek(0)
    text = ""
    pdf_document = fitz.open(
        stream=pdf_file.read(),
        filetype="pdf"
    )
    for page in pdf_document:
        text += page.get_text()
    pdf_document.close()
    return text

# -----------------------------
# Extract Text From DOCX
# -----------------------------
def extract_text_from_docx(docx_file):
    docx_file.seek(0)
    doc = docx.Document(docx_file)
    full_text = []
    for para in doc.paragraphs:
        full_text.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                full_text.append(cell.text)
    return "\n".join(full_text)

# -----------------------------
# Unified Text Extractor
# -----------------------------
def extract_text(file):
    name = file.name.lower()
    try:
        if name.endswith(".pdf"):
            return extract_text_from_pdf(file)
        elif name.endswith(".docx"):
            return extract_text_from_docx(file)
        else:
            raise ValueError("Unsupported file format. Please upload PDF or DOCX.")
    except Exception as e:
        st.error(f"Error reading file: {str(e)}")
        return None

# -----------------------------
# Extract Skills using spaCy
# -----------------------------
def extract_skills(text, common_skills):
    nlp = load_spacy_model()
    # Process text
    doc = nlp(text.lower())
    
    # Initialize PhraseMatcher
    matcher = PhraseMatcher(nlp.vocab, attr="LOWER")
    patterns = [nlp.make_doc(skill.lower()) for skill in common_skills]
    matcher.add("SKILLS", patterns)
    
    matches = matcher(doc)
    found_skills = set()
    for match_id, start, end in matches:
        span = doc[start:end]
        found_skills.add(span.text)
        
    return sorted(list(found_skills))

# -----------------------------
# Analyze Resume
# -----------------------------
def analyze_resume(resume_text, job_description):
    resume_skills = extract_skills(resume_text, COMMON_SKILLS)
    jd_skills = extract_skills(job_description, COMMON_SKILLS)
    
    matched_skills = sorted(
        list(set(resume_skills) & set(jd_skills))
    )
    
    missing_skills = sorted(
        list(set(jd_skills) - set(resume_skills))
    )
    
    if len(jd_skills) > 0:
        score = int(
            (len(matched_skills) / len(jd_skills)) * 100
        )
    else:
        score = 0
        
    suggestions = []
    
    if score >= 80:
        suggestions.append(
            "Excellent resume match for the job description."
        )
    elif score >= 60:
        suggestions.append(
            "Good match. Add more relevant skills to improve ATS score."
        )
    else:
        suggestions.append(
            "Low ATS score. Resume needs improvement."
        )
        
    if missing_skills:
        suggestions.append(
            f"Include missing skills if you have experience with them: {', '.join(missing_skills)}"
        )
        
    return {
        "score": score,
        "matched": matched_skills,
        "missing": missing_skills,
        "suggestions": suggestions
    }

# -----------------------------
# Streamlit Configuration & CSS
# -----------------------------
st.set_page_config(
    page_title="AI Resume Analyzer",
    layout="wide"
)

# Inject modern, premium styling
st.markdown("""
    <style>
    .stApp {
        background-color: #F8FAFC;
    }
    .main-header {
        background: linear-gradient(135deg, #1E3A8A 0%, #3B82F6 100%);
        padding: 30px;
        border-radius: 12px;
        color: white;
        margin-bottom: 25px;
        box-shadow: 0 4px 6px -1px rgb(0 0 0 / 0.1), 0 2px 4px -2px rgb(0 0 0 / 0.1);
    }
    .card {
        background-color: white;
        padding: 20px;
        border-radius: 8px;
        box-shadow: 0 1px 3px 0 rgb(0 0 0 / 0.1), 0 1px 2px -1px rgb(0 0 0 / 0.1);
        border: 1px solid #E2E8F0;
        margin-bottom: 15px;
    }
    .stButton>button {
        background-color: #2563EB;
        color: white;
        border-radius: 8px;
        border: none;
        font-weight: 600;
        padding: 10px 24px;
        transition: all 0.2s ease-in-out;
    }
    .stButton>button:hover {
        background-color: #1D4ED8;
        transform: translateY(-1px);
    }
    </style>
""", unsafe_allow_html=True)

# Initialize Session State
if 'user' not in st.session_state:
    st.session_state['user'] = None

# -----------------------------
# Navigation Sidebar
# -----------------------------
st.sidebar.markdown("""
    <div style="text-align: center; padding: 10px 0;">
        <h2 style="margin: 0; color: #1E3A8A;">📄 ATS System</h2>
        <p style="color: #64748B; font-size: 14px;">Next-Gen Resume Analysis</p>
    </div>
""", unsafe_allow_html=True)

if st.session_state['user']:
    st.sidebar.markdown(f"""
        <div style="background-color: #EFF6FF; padding: 15px; border-radius: 8px; margin-bottom: 15px; border: 1px solid #BFDBFE;">
            <p style="margin: 0; font-size: 11px; color: #1E40AF; font-weight: bold; letter-spacing: 0.05em;">ACTIVE USER</p>
            <p style="margin: 5px 0 0 0; font-size: 16px; font-weight: bold; color: #1E3A8A;">{st.session_state['user']['fullname']}</p>
            <p style="margin: 2px 0 0 0; font-size: 12px; color: #64748B;">@{st.session_state['user']['username']}</p>
        </div>
    """, unsafe_allow_html=True)
    
    page = st.sidebar.radio("Navigation", ["🔍 Analyze Resume", "📋 Saved Reports", "👤 Profile Management", "🚪 Logout"])
else:
    st.sidebar.markdown("""
        <div style="background-color: #F1F5F9; padding: 15px; border-radius: 8px; margin-bottom: 15px; border: 1px solid #E2E8F0;">
            <p style="margin: 0; font-size: 14px; font-weight: bold; color: #475569;">Guest Mode</p>
            <p style="margin: 5px 0 0 0; font-size: 12px; color: #64748B;">Log in to save report logs and manage a customized skills profile.</p>
        </div>
    """, unsafe_allow_html=True)
    
    page = st.sidebar.radio("Navigation", ["🔍 Analyze Resume", "🔑 Login / Signup"])

# -----------------------------
# PAGE: Analyze Resume
# -----------------------------
if page == "🔍 Analyze Resume":
    st.markdown("""
        <div class="main-header">
            <h1 style="color: white; margin: 0; font-size: 28px;">📄 AI Resume Analyzer</h1>
            <p style="color: #93C5FD; margin: 5px 0 0 0; font-size: 16px;">Intelligent similarity scoring using spaCy token matching</p>
        </div>
    """, unsafe_allow_html=True)
    
    col1, col2 = st.columns(2)
    
    with col1:
        st.subheader("1. Upload Resume")
        uploaded_file = st.file_uploader(
            "Upload PDF or DOCX Resume",
            type=["pdf", "docx"]
        )
        
    with col2:
        st.subheader("2. Job Description")
        job_description = st.text_area(
            "Enter Job Description Requirements",
            height=200,
            placeholder="Search matching words by pasting the job listing requirements here..."
        )
        
    analyze_btn = st.button("🔍 Run Analysis", use_container_width=True)
    
    if analyze_btn:
        if uploaded_file is None:
            st.warning("Please upload a PDF or DOCX resume.")
            st.stop()
            
        if not job_description.strip():
            st.warning("Please enter a job description.")
            st.stop()
            
        with st.spinner("Extracting text contents..."):
            resume_text = extract_text(uploaded_file)
            
        if resume_text is None:
            st.stop()
            
        with st.spinner("Analyzing matching elements..."):
            result = analyze_resume(resume_text, job_description)
            
        st.markdown("<hr/>", unsafe_allow_html=True)
        st.subheader("📊 ATS Analysis Report")
        
        score = result['score']
        if score >= 80:
            score_color = "#10B981"
        elif score >= 60:
            score_color = "#F59E0B"
        else:
            score_color = "#EF4444"
            
        # Display Score Card
        c1, c2, c3 = st.columns([1, 2, 1])
        with c2:
            st.markdown(f"""
                <div style="background-color: white; padding: 25px; border-radius: 12px; border-left: 6px solid {score_color}; box-shadow: 0 4px 6px -1px rgb(0 0 0 / 0.1); text-align: center;">
                    <span style="font-size: 14px; color: #64748B; font-weight: bold; text-transform: uppercase; letter-spacing: 0.05em;">ATS Match Score</span>
                    <h2 style="font-size: 56px; margin: 10px 0; color: {score_color}; font-weight: 800;">{score}%</h2>
                    <p style="margin: 0; color: #475569; font-size: 14px;">Calculated using NLP tokenizer skill matches</p>
                </div>
            """, unsafe_allow_html=True)
            
        st.markdown("<br/>", unsafe_allow_html=True)
        st.progress(score / 100)
        
        # Save report database records if logged in
        if st.session_state['user']:
            database.save_report(
                user_id=st.session_state['user']['id'],
                resume_name=uploaded_file.name,
                job_description=job_description,
                score=score,
                matched_skills=result['matched'],
                missing_skills=result['missing'],
                suggestions=result['suggestions']
            )
            st.success("💾 Analysis report saved successfully to your Profile History!")
        else:
            st.info("💡 **Tip**: Log in to automatically log report history and track progress over time.")
            
        # Visual Grid for Skills
        col_m1, col_m2 = st.columns(2)
        with col_m1:
            st.markdown(f"#### ✅ Matched Skills ({len(result['matched'])})")
            if result["matched"]:
                badges = "".join([f'<span style="display: inline-block; background-color: #DEF7EC; color: #03543F; font-size: 12px; font-weight: 600; padding: 5px 10px; border-radius: 9999px; margin: 4px; border: 1px solid #BCF0DA;">{skill}</span>' for skill in result['matched']])
                st.markdown(f'<div style="margin-bottom: 15px;">{badges}</div>', unsafe_allow_html=True)
            else:
                st.warning("No matched skills identified.")
                
        with col_m2:
            st.markdown(f"#### ❌ Missing Skills ({len(result['missing'])})")
            if result["missing"]:
                badges = "".join([f'<span style="display: inline-block; background-color: #FDE8E8; color: #9B1C1C; font-size: 12px; font-weight: 600; padding: 5px 10px; border-radius: 9999px; margin: 4px; border: 1px solid #FBD5D5;">{skill}</span>' for skill in result['missing']])
                st.markdown(f'<div style="margin-bottom: 15px;">{badges}</div>', unsafe_allow_html=True)
            else:
                st.success("No missing skills from the job description!")
                
        # Display suggestions
        st.subheader("💡 Suggestions")
        for suggestion in result["suggestions"]:
            st.info(suggestion)
            
        # PDF Generation & Download
        with st.spinner("Compiling PDF report..."):
            pdf_data = pdf_generator.generate_pdf_report(
                result=result,
                resume_name=uploaded_file.name,
                username=st.session_state['user']['fullname'] if st.session_state['user'] else "Guest"
            )
            
        st.download_button(
            label="📥 Download PDF Report",
            data=pdf_data,
            file_name=f"resume_analysis_{int(datetime.now().timestamp())}.pdf",
            mime="application/pdf",
            use_container_width=True
        )

# -----------------------------
# PAGE: Saved Reports
# -----------------------------
elif page == "📋 Saved Reports":
    st.markdown("""
        <div class="main-header">
            <h1 style="color: white; margin: 0; font-size: 28px;">📋 Saved Reports History</h1>
            <p style="color: #93C5FD; margin: 5px 0 0 0; font-size: 16px;">View past resume analytics and export files</p>
        </div>
    """, unsafe_allow_html=True)
    
    if 'delete_success' in st.session_state and st.session_state['delete_success']:
        st.success(st.session_state['delete_success'])
        del st.session_state['delete_success']
        
    reports = database.get_user_reports(st.session_state['user']['id'])
    
    if not reports:
        st.info("No saved report history. Head to 'Analyze Resume' to create a new report.")
    else:
        for report in reports:
            score = report['score']
            score_color = '#10B981' if score >= 80 else ('#F59E0B' if score >= 60 else '#EF4444')
            
            with st.container():
                st.markdown(f"""
                    <div style="background-color: white; padding: 20px; border-radius: 8px; border-left: 5px solid {score_color}; box-shadow: 0 1px 3px rgb(0 0 0 / 0.1); margin-bottom: 15px;">
                        <div style="display: flex; justify-content: space-between; align-items: center;">
                            <div>
                                <h4 style="margin: 0; color: #1F2937;">📄 {report['resume_name']}</h4>
                                <span style="font-size: 12px; color: #6B7280;">Analyzed on: {report['created_at']}</span>
                            </div>
                            <div style="text-align: right;">
                                <span style="font-size: 24px; font-weight: bold; color: {score_color};">{score}%</span>
                            </div>
                        </div>
                    </div>
                """, unsafe_allow_html=True)
                
                with st.expander("Details"):
                    col_m1, col_m2 = st.columns(2)
                    with col_m1:
                        st.markdown("**Matched Skills**")
                        if report['matched_skills']:
                            st.write(", ".join(report['matched_skills']))
                        else:
                            st.write("None")
                    with col_m2:
                        st.markdown("**Missing Skills**")
                        if report['missing_skills']:
                            st.write(", ".join(report['missing_skills']))
                        else:
                            st.write("None")
                            
                    st.markdown("**Suggestions**")
                    for suggestion in report['suggestions']:
                        st.info(suggestion)
                        
                    # Download / Delete Actions
                    c_dl, c_del = st.columns(2)
                    with c_dl:
                        pdf_data = pdf_generator.generate_pdf_report(
                            result={
                                'score': report['score'],
                                'matched': report['matched_skills'],
                                'missing': report['missing_skills'],
                                'suggestions': report['suggestions']
                            },
                            resume_name=report['resume_name'],
                            username=st.session_state['user']['fullname']
                        )
                        st.download_button(
                            label="📥 Download PDF",
                            data=pdf_data,
                            file_name=f"report_{report['id']}.pdf",
                            mime="application/pdf",
                            key=f"dl_{report['id']}"
                        )
                    with c_del:
                        if st.button("🗑️ Delete Record", key=f"del_{report['id']}"):
                            database.delete_report(report['id'], st.session_state['user']['id'])
                            st.session_state['delete_success'] = "Report deleted successfully!"
                            st.rerun()

# -----------------------------
# PAGE: Profile Management
# -----------------------------
elif page == "👤 Profile Management":
    st.markdown("""
        <div class="main-header">
            <h1 style="color: white; margin: 0; font-size: 28px;">👤 Profile Management</h1>
            <p style="color: #93C5FD; margin: 5px 0 0 0; font-size: 16px;">Manage your developer profile and targeted skillset list</p>
        </div>
    """, unsafe_allow_html=True)
    
    user_id = st.session_state['user']['id']
    profile = database.get_user_profile(user_id)
    
    if profile:
        fullname = st.text_input("Full Name", value=profile['fullname'])
        email = st.text_input("Email Address", value=profile['email'])
        profile_skills = st.text_area("Target Skills (comma-separated list)", value=profile['profile_skills'], help="e.g. Python, SQL, React")
        profile_summary = st.text_area("Professional Summary", value=profile['profile_summary'], height=150)
        
        if st.button("Update Profile Information", use_container_width=True):
            database.update_user_profile(user_id, email, fullname, profile_skills, profile_summary)
            st.session_state['user'] = database.get_user_profile(user_id)
            st.success("Profile updated successfully!")
            st.rerun()

# -----------------------------
# PAGE: Login / Signup
# -----------------------------
elif page == "🔑 Login / Signup":
    st.markdown("""
        <div class="main-header">
            <h1 style="color: white; margin: 0; font-size: 28px;">🔑 User Portal</h1>
            <p style="color: #93C5FD; margin: 5px 0 0 0; font-size: 16px;">Log in or create a profile to persist analyses</p>
        </div>
    """, unsafe_allow_html=True)
    
    auth_tabs = st.tabs(["🔒 Secure Login", "📝 Create Account"])
    
    with auth_tabs[0]:
        st.subheader("Login Credentials")
        login_user = st.text_input("Username", key="login_user")
        login_pwd = st.text_input("Password", type="password", key="login_pwd")
        
        if st.button("Log In", use_container_width=True):
            if not login_user.strip() or not login_pwd.strip():
                st.error("Please fill in both fields.")
            else:
                user = database.authenticate_user(login_user, login_pwd)
                if user:
                    st.session_state['user'] = user
                    st.success("Successfully logged in!")
                    st.rerun()
                else:
                    st.error("Invalid username or password.")
                    
    with auth_tabs[1]:
        st.subheader("Register a Profile")
        reg_fullname = st.text_input("Full Name", key="reg_fullname")
        reg_email = st.text_input("Email Address", key="reg_email")
        reg_user = st.text_input("Choose Username", key="reg_user")
        reg_pwd = st.text_input("Choose Password", type="password", key="reg_pwd")
        
        if st.button("Register Account", use_container_width=True):
            if not reg_fullname.strip() or not reg_email.strip() or not reg_user.strip() or not reg_pwd.strip():
                st.error("Please fill in all fields.")
            elif len(reg_pwd) < 6:
                st.error("Password must be at least 6 characters.")
            else:
                success = database.create_user(reg_user, reg_pwd, reg_email, reg_fullname)
                if success:
                    st.success("Registration successful! You can now log in using the tab above.")
                else:
                    st.error("Username already exists. Please choose a different username.")

# -----------------------------
# PAGE: Logout Handling
# -----------------------------
elif page == "🚪 Logout":
    st.session_state['user'] = None
    st.success("Logged out successfully.")
    st.rerun()